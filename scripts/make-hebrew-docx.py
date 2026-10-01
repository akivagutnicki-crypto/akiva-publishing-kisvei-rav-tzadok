#!/usr/bin/env python3
"""Build a right-to-left Hebrew .docx from transcribed page files.

Usage: scripts/make-hebrew-docx.py <pages dir> <out.docx> "<title>" ["<subtitle>"]

Each page file (NNN.txt, in page order) holds the text of one printed page:
  "# ..."      a heading: "מצות ..." -> Heading 1, "פרק ..." -> Heading 2, else Heading 3
  "~ ..."      the page opens in the middle of the previous page's paragraph
  blank line   a paragraph break
  **...**      bold (the opening words of a paragraph, as in the print)
  "! name.png" a figure, read from <pages dir>/figs/ (scanned at 300 dpi)
  "| a | b"    a table row (a block of such lines is one table; the first row is the header)
Files whose names start with "_" (notes) are skipped.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image

FONT = 'David'
NIQQUD = re.compile(r'[֑-ׇ]')


def blocks(pages_dir):
    """[(kind, text)] with paragraphs joined across pages."""
    out = []
    for f in sorted(p for p in Path(pages_dir).glob('*.txt') if not p.name.startswith('_')):
        for k, chunk in enumerate(re.split(r'\n\s*\n', f.read_text(encoding='utf-8').strip())):
            for line in [l for l in chunk.split('\n') if l.startswith('# ')]:
                out.append(('h', line[2:].strip()))
            lines = [l.strip() for l in chunk.split('\n') if l.strip() and not l.startswith('# ')]
            if lines and lines[0].startswith('! '):
                out.append(('img', str(Path(pages_dir) / 'figs' / lines[0][2:].strip())))
                continue
            if lines and all(l.startswith('|') for l in lines):
                out.append(('table', [[c.strip() for c in l.strip('|').split('|')] for l in lines]))
                continue
            text = ' '.join(l.strip() for l in chunk.split('\n') if l.strip() and not l.startswith('# '))
            if not text:
                continue
            if text.startswith("~ ") and k == 0 and out and out[-1][0] == 'p':
                out[-1] = ('p', out[-1][1] + ' ' + text[2:])
            else:
                out.append(('p', text.removeprefix('~ ')))
    return out


def rtl(paragraph, size=14, bold=None):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.append(OxmlElement('w:bidi'))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if bold is None else WD_ALIGN_PARAGRAPH.CENTER


def add_runs(paragraph, text, size=14, bold_all=False, color=None):
    for part in re.split(r'(\*\*[^*]+\*\*)', text):
        if not part:
            continue
        bold = bold_all or part.startswith('**')
        run = paragraph.add_run(part.strip('*') if part.startswith('**') else part)
        rPr = run._r.get_or_add_rPr()
        # children in the order the OOXML schema requires (Word rejects other orders)
        fonts = OxmlElement('w:rFonts')
        for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
            fonts.set(qn(a), FONT)
        children = [fonts]
        if bold:
            children += [OxmlElement('w:b'), OxmlElement('w:bCs')]
        if color:
            c = OxmlElement('w:color'); c.set(qn('w:val'), color); children.append(c)
        for tag in ('w:sz', 'w:szCs'):
            e = OxmlElement(tag); e.set(qn('w:val'), str(size * 2)); children.append(e)
        children.append(OxmlElement('w:rtl'))
        for e in children:
            rPr.append(e)


def heading_level(text):
    plain = NIQQUD.sub('', text)
    if re.match(r'^(מצות|מצוה|מצוות|מ"ע|מל"ת)', plain.replace('״', '"')):
        return 1
    if plain.startswith('פרק'):
        return 2
    return 3


def build(pages_dir, out, title, subtitle=None):
    doc = Document()
    sec = doc.sections[0]
    sectPr = sec._sectPr
    grid = sectPr.find(qn('w:docGrid'))
    (grid.addprevious if grid is not None else sectPr.append)(OxmlElement('w:bidi'))
    p = doc.add_paragraph(style='Title'); rtl(p, bold=True); add_runs(p, title, size=26, bold_all=True)
    if subtitle:
        p = doc.add_paragraph(); rtl(p, bold=True); add_runs(p, subtitle, size=13, color='555555')
    for kind, text in blocks(pages_dir):
        if kind == 'h':
            level = heading_level(text)
            p = doc.add_paragraph(style=f'Heading {level}'); rtl(p, bold=True)
            add_runs(p, text, size={1: 20, 2: 17, 3: 15}[level], bold_all=True, color='17372D')
        elif kind == 'img':
            width = min(Image.open(text).size[0] / 300, 6.0)
            doc.add_picture(text, width=Inches(width))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind == 'table':
            table = doc.add_table(rows=len(text), cols=len(text[0]))
            table.style = 'Table Grid'
            tblPr = table._tbl.tblPr
            style = tblPr.find(qn('w:tblStyle'))
            bidi = OxmlElement('w:bidiVisual')  # columns run right to left
            if style is not None:
                style.addnext(bidi)
            else:
                tblPr.insert(0, bidi)
            for r, row in enumerate(text):
                for c, val in enumerate(row):
                    p = table.cell(r, c).paragraphs[0]; rtl(p, bold=True)
                    add_runs(p, val, size=12, bold_all=(r == 0))
            doc.add_paragraph()
        else:
            p = doc.add_paragraph(); rtl(p)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.35
            add_runs(p, text)
    doc.save(out)
    return out


if __name__ == '__main__':
    print(build(*sys.argv[1:]))
